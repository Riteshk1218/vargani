"""Every list can be downloaded as CSV or Excel by the people the Organization Admin allows."""

import io

import pytest
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.organizations.models import OrganizationMembership
from apps.payments.tests.test_payments import collect, month_line, setup  # noqa: F401 - fixture

pytestmark = pytest.mark.django_db

LISTS = [
    "contributions",
    "payments",
    "receipts",
    "expenses",
    "other-income",
    "people",
    "events",
    "audit-logs",
    "members",
]


def _member_id(admin, user):
    return next(m["id"] for m in admin.get("/api/v1/members/").data["results"] if m["user"]["id"] == str(user.pk))


def test_admin_downloads_every_list_as_csv_and_excel(setup):  # noqa: F811
    s = setup
    assert collect(s.admin, s, [month_line(s, "2027-04-01", "500")]).status_code == 201
    for name in LISTS:
        response = s.admin.get(f"/api/v1/{name}/download/", {"file": "csv"})
        assert response.status_code == 200, (name, response.content[:300])
        assert response["Content-Type"].startswith("text/csv") and "attachment" in response["Content-Disposition"]
    csv_text = s.admin.get("/api/v1/contributions/download/", {"event": str(s.event.pk)}).content.decode("utf-8-sig")
    assert "Ritesh Patil" in csv_text and "500.00" in csv_text and "Monthly Vargani" in csv_text

    response = s.admin.get("/api/v1/payments/download/", {"file": "xlsx", "event": str(s.event.pk)})
    sheet = load_workbook(io.BytesIO(response.content)).active
    values = [cell for row in sheet.iter_rows(values_only=True) for cell in row]
    assert "Ritesh Patil" in values and 500 in values
    assert AuditLog.objects.filter(organization=s.org, action="export", entity_type="list.payment").exists()


def test_search_and_filters_apply_to_the_file(setup):  # noqa: F811
    s = setup
    text = s.admin.get("/api/v1/people/download/", {"search": "nobody-like-this"}).content.decode("utf-8-sig")
    assert "Ritesh Patil" not in text
    assert "Ritesh Patil" in s.admin.get("/api/v1/people/download/").content.decode("utf-8-sig")


def test_admin_decides_who_may_download(setup, add_member, client_for):  # noqa: F811
    s = setup
    assert collect(s.admin, s, [month_line(s, "2027-04-01", "500")]).status_code == 201
    treasurer_user = add_member(s.org, "treasurer")
    collector_user = add_member(s.org, "collector")
    treasurer = client_for(treasurer_user, s.org)
    collector = client_for(collector_user, s.org)

    # Role defaults: the treasurer may download, a collector may not.
    assert treasurer.get("/api/v1/payments/download/").status_code == 200
    assert collector.get("/api/v1/payments/download/").status_code == 403

    # The admin switches it off for the treasurer: lists and report exports both refuse.
    response = s.admin.patch(f"/api/v1/members/{_member_id(s.admin, treasurer_user)}/", {"can_download": False})
    assert response.status_code == 200 and response.data["can_download"] is False
    assert treasurer.get("/api/v1/payments/download/").status_code == 403
    assert "reports.export" not in treasurer.get(f"/api/v1/me/organizations/{s.org.pk}/").data["permissions"]
    assert treasurer.post("/api/v1/reports/contributions/exports/", {"format": "csv"}, format="json").status_code == 403

    # ...and on for the collector, who still only gets their own records.
    s.admin.patch(f"/api/v1/members/{_member_id(s.admin, collector_user)}/", {"can_download": True})
    response = collector.get("/api/v1/payments/download/")
    assert response.status_code == 200 and "Ritesh Patil" not in response.content.decode("utf-8-sig")
    assert AuditLog.objects.filter(organization=s.org, action="role_change", changed_fields__contains=["can_download"])

    # Back to the role's default.
    s.admin.patch(f"/api/v1/members/{_member_id(s.admin, treasurer_user)}/", {"can_download": None}, format="json")
    assert OrganizationMembership.objects.get(user=treasurer_user).can_download is None
    assert treasurer.get("/api/v1/payments/download/").status_code == 200


def test_only_admins_change_download_access(setup, add_member, client_for):  # noqa: F811
    s = setup
    treasurer_user = add_member(s.org, "treasurer")
    treasurer = client_for(treasurer_user, s.org)
    member_id = _member_id(s.admin, treasurer_user)
    assert treasurer.patch(f"/api/v1/members/{member_id}/", {"can_download": True}).status_code == 403


def test_unknown_file_type_is_refused(setup):  # noqa: F811
    assert setup.admin.get("/api/v1/people/download/", {"file": "pdf"}).status_code == 400
