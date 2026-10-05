"""Render a Report as JSON-ready data, CSV, Excel or PDF."""

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from django.utils import timezone
from django.utils.translation import gettext as _

from apps.core.money import indian_format
from apps.core.pdf import render_pdf, storage_data_uri

from .registry import Report

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
PDF_MAX_ROWS = 5000


def json_value(value):
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, datetime | date):
        return value.isoformat()
    return value


def as_json(report: Report, limit: int) -> dict:
    return {
        "title": report.title,
        "columns": [{"key": c.key, "label": c.label, "kind": c.kind} for c in report.columns],
        "rows": [{k: json_value(v) for k, v in row.items()} for row in report.rows[:limit]],
        "row_count": len(report.rows),
        "truncated": len(report.rows) > limit,
        "totals": {k: json_value(v) for k, v in report.totals.items()},
        "notes": report.notes,
    }


def _text(value, kind):
    if value is None:
        return ""
    if kind == "money" and isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _header_lines(report, meta):
    return [report.title, meta["organization"], meta["scope"], _("Generated on %(when)s") % {"when": meta["generated"]}]


def to_csv(report: Report, meta: dict) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for line in _header_lines(report, meta):
        writer.writerow([line])
    writer.writerow([])
    writer.writerow([c.label for c in report.columns])
    for row in report.rows:
        writer.writerow([_safe_csv(_text(row.get(c.key), c.kind)) for c in report.columns])
    if report.totals:
        writer.writerow([])
        for key, value in report.totals.items():
            writer.writerow([_total_label(report, key), _text(value, "money" if isinstance(value, Decimal) else "")])
    # BOM so Excel opens Marathi/Hindi text as UTF-8.
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def _safe_csv(text: str) -> str:
    """Neutralise spreadsheet formulas in user data (CSV injection)."""
    if text and text[0] in "=+-@\t\r" and not _is_number(text):
        return "'" + text
    return text


def _is_number(text):
    try:
        Decimal(text)
        return True
    except Exception:  # noqa: BLE001
        return False


def _total_label(report, key):
    for column in report.columns:
        if column.key == key:
            return _("Total: %(label)s") % {"label": column.label}
    return {"count": _("Rows"), "balance": _("Balance")}.get(key, key)


def to_xlsx(report: Report, meta: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31]
    for line in _header_lines(report, meta):
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([c.label for c in report.columns])
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FDE7D7")
    widths = [len(c.label) for c in report.columns]
    for row in report.rows:
        values = []
        for i, c in enumerate(report.columns):
            value = row.get(c.key)
            if isinstance(value, str):
                value = _safe_csv(value)
            values.append(value)
            widths[i] = max(widths[i], min(len(_text(value, c.kind)), 60))
        ws.append(values)
        for i, c in enumerate(report.columns):
            cell = ws.cell(row=ws.max_row, column=i + 1)
            if c.kind == "money":
                cell.number_format = "#,##0.00"
            elif c.kind == "date":
                cell.number_format = "yyyy-mm-dd"
            elif c.kind == "datetime":
                cell.number_format = "yyyy-mm-dd hh:mm"
    if report.totals:
        ws.append([])
        for key, value in report.totals.items():
            ws.append([_total_label(report, key), value])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            if isinstance(value, Decimal):
                ws.cell(row=ws.max_row, column=2).number_format = "#,##0.00"
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, width + 2)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for cell in ws[header_row]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pdf_value(value, kind):
    if value is None or value == "":
        return ""
    if kind == "money":
        return "₹" + indian_format(value)
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    return value


def to_pdf(report: Report, meta: dict) -> bytes:
    rows = [
        [{"value": _pdf_value(row.get(c.key), c.kind), "kind": c.kind} for c in report.columns]
        for row in report.rows[:PDF_MAX_ROWS]
    ]
    totals = [
        {"label": _total_label(report, k), "value": _pdf_value(v, "money" if isinstance(v, Decimal) else "")}
        for k, v in report.totals.items()
    ]
    return render_pdf(
        "reports/report_pdf.html",
        {
            "report": report,
            "rows": rows,
            "totals": totals,
            "meta": meta,
            "logo": storage_data_uri(meta.get("logo")),
            "landscape": len(report.columns) > 6,
            "truncated": len(report.rows) > PDF_MAX_ROWS,
            "pdf_max_rows": PDF_MAX_ROWS,
        },
    )


RENDERERS = {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf}


def meta_for(org, scope_label: str) -> dict:
    return {
        "organization": org.name,
        "logo": org.logo or None,
        "scope": scope_label,
        "generated": timezone.localtime().strftime("%d-%m-%Y %H:%M"),
    }
