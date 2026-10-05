import csv
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.events.models import Event
from apps.expenses.models import ExpenseCategory
from apps.organizations.models import FinancialYear
from apps.payments.tests.test_payments import collect, month_line, setup  # noqa: F401 - fixture
from apps.people.models import Person

pytestmark = pytest.mark.django_db
TODAY = date.today().isoformat()


def spend(s, amount, category="Decoration", approve=True):
    cat = ExpenseCategory.objects.get(organization=s.org, name=category)
    payload = {
        "event": str(s.event.pk),
        "category": str(cat.pk),
        "amount": amount,
        "expense_date": TODAY,
        "payee": "Vendor",
    }
    expense_id = s.admin.post("/api/v1/expenses/", payload, format="json").data["id"]
    if approve:
        s.admin.post(f"/api/v1/expenses/{expense_id}/submit/")
        s.admin.post(f"/api/v1/expenses/{expense_id}/approve/")
    return expense_id


def income(s, amount, source="donation_box"):
    payload = {"event": str(s.event.pk), "source": source, "amount": amount, "received_on": TODAY}
    return s.admin.post("/api/v1/other-income/", payload, format="json").data["id"]


@pytest.fixture
def books(setup):  # noqa: F811
    """A hand-checked set of books.

    Contributions confirmed: 500.25 + 700.50 + 1001.00 = 2201.75 (a cancelled 999.99 pending cheque is excluded)
    Other income recorded:   1234.56 + 0.44 = 1235.00 (a voided 100.00 is excluded)
    Expenses approved:       2500.10 + 300.20 = 2800.30 (a draft 50.00 and a voided 75.00 are excluded)
    Income  = 2201.75 + 1235.00 = 3436.75
    Balance = 3436.75 − 2800.30 = 636.45
    """
    s = setup
    collect(s.admin, s, [month_line(s, "2027-04-01", "500.25")])
    collect(s.admin, s, [month_line(s, "2027-05-01", "700.50")], mode="upi", reference_number="UPI1")
    collect(s.admin, s, [{"contribution_type": str(s.donation.pk), "amount": "1001.00"}])
    pending = collect(
        s.admin,
        s,
        [{"contribution_type": str(s.donation.pk), "amount": "999.99"}],
        mode="cheque",
        reference_number="000123",
        status="pending",
    ).data
    s.admin.post(f"/api/v1/payments/{pending['id']}/cancel/", {"reason": "Bounced"}, format="json")
    income(s, "1234.56")
    income(s, "0.44", source="interest")
    voided = income(s, "100.00", source="sale")
    s.admin.post(f"/api/v1/other-income/{voided}/void/", {"reason": "Wrong"}, format="json")
    spend(s, "2500.10")
    spend(s, "300.20", category="Sound")
    spend(s, "50.00", approve=False)
    void_me = spend(s, "75.00")
    s.admin.post(f"/api/v1/expenses/{void_me}/void/", {"reason": "Duplicate"}, format="json")
    return s


def report(client, s, kind, **params):
    response = client.get(f"/api/v1/reports/{kind}/", {"event": str(s.event.pk), **params})
    assert response.status_code == 200, response.data
    return response.data


def test_income_expense_and_balance_match_to_the_paisa(books):
    s = books
    data = report(s.admin, s, "income-expense")
    assert data["totals"] == {"income": "3436.75", "expense": "2800.30", "balance": "636.45"}
    heads = {(r["section"], r["head"]): (r["income"] or r["expense"]) for r in data["rows"]}
    assert heads[("Income", "Monthly Vargani")] == "1200.75"
    assert heads[("Income", "Donation")] == "1001.00"
    assert heads[("Income", "Donation box")] == "1234.56"
    assert heads[("Income", "Bank interest")] == "0.44"
    assert heads[("Expense", "Decoration")] == "2500.10"
    assert heads[("Expense", "Sound")] == "300.20"

    dashboard = s.admin.get("/api/v1/dashboard/summary/", {"event": str(s.event.pk)}).data["financial"]
    assert dashboard["collected"] == "2201.75"
    assert dashboard["other_income"] == "1235.00"
    assert dashboard["income"] == "3436.75"
    assert dashboard["expenses"] == "2800.30"
    assert dashboard["balance"] == "636.45"
    assert dashboard["pending_expenses"] == 0

    event_wise = report(s.admin, s, "event-wise", financial_year=str(s.fy.pk))
    row = next(r for r in event_wise["rows"] if r["event"] == s.event.name)
    assert (row["contributions"], row["other_income"], row["expenses"], row["balance"]) == (
        "2201.75",
        "1235.00",
        "2800.30",
        "636.45",
    )


def test_collection_reports(books):
    s = books
    contributions = report(s.admin, s, "contributions")
    assert contributions["totals"]["amount"] == "2201.75" and contributions["row_count"] == 3
    assert all(r["receipt"].startswith("GC-") for r in contributions["rows"])
    modes = {r["mode"]: r["amount"] for r in report(s.admin, s, "payment-mode")["rows"]}
    assert modes == {"Cash": "1501.25", "UPI": "700.50"}
    assert report(s.admin, s, "collector-wise")["totals"]["amount"] == "2201.75"
    funds = report(s.admin, s, "fund-wise")
    assert funds["totals"] == {"contributions": "2201.75", "other_income": "1235.00", "total": "3436.75"}
    months = report(s.admin, s, "monthly-collection")
    assert months["totals"]["amount"] == "2201.75"
    receipts = report(s.admin, s, "receipts")
    assert receipts["row_count"] == 3 and receipts["totals"]["amount"] == "2201.75"
    expenses = report(s.admin, s, "expenses")
    assert expenses["totals"]["amount"] == "2800.30" and expenses["row_count"] == 2
    assert report(s.admin, s, "expenses", status="all")["row_count"] == 4
    assert report(s.admin, s, "other-income")["totals"]["amount"] == "1235.00"
    # Date filters.
    assert report(s.admin, s, "contributions", date_from="2000-01-01", date_to="2000-01-02")["row_count"] == 0
    bad = s.admin.get("/api/v1/reports/contributions/", {"date_from": "2027-02-01", "date_to": "2027-01-01"})
    assert bad.status_code == 400


def test_dues_report(setup):  # noqa: F811
    s = setup
    s.monthly.is_mandatory = True
    s.monthly.target_amount = Decimal("500")
    s.monthly.start_date = date.today().replace(day=1)
    s.monthly.save()
    from apps.events.models import EventParticipant

    EventParticipant.objects.create(organization=s.org, event=s.event, person=s.ritesh)
    collect(s.admin, s, [month_line(s, date.today().replace(day=1).isoformat(), "200")])
    data = report(s.admin, s, "dues")
    assert data["rows"][0]["person"] == "Ritesh Patil" and data["rows"][0]["due"] == "300.00"
    assert s.admin.get("/api/v1/reports/dues/").status_code == 400  # needs an event


def test_general_reports_and_list(setup, add_member, client_for):  # noqa: F811
    s = setup
    Person.objects.create(organization=s.org, full_name="Asha Kale", mobile="+919822001122")
    people = s.admin.get("/api/v1/reports/people/").data
    assert [r["name"] for r in people["rows"]] == ["Asha Kale", "Ritesh Patil"]
    events = s.admin.get("/api/v1/reports/events/").data
    assert events["rows"][0]["name"] == s.event.name
    listing = s.admin.get("/api/v1/reports/").data
    assert {r["type"] for r in listing} >= {"people", "income-expense", "audit"}
    assert s.admin.get("/api/v1/reports/nope/").status_code == 404


def test_report_permissions(books, add_member, client_for):
    s = books
    secretary = client_for(add_member(s.org, "secretary"), s.org)
    # Secretary reads financial reports but exports only general ones.
    assert secretary.get("/api/v1/reports/income-expense/", {"event": str(s.event.pk)}).status_code == 200
    assert secretary.post("/api/v1/reports/people/exports/", {"format": "csv"}, format="json").status_code == 202
    denied = secretary.post("/api/v1/reports/income-expense/exports/", {"format": "csv"}, format="json")
    assert denied.status_code == 403
    assert secretary.get("/api/v1/reports/audit/").status_code == 403
    types = {r["type"] for r in secretary.get("/api/v1/reports/").data}
    assert "audit" not in types and "income-expense" in types
    for role in ("collector", "volunteer", "member"):
        client = client_for(add_member(s.org, role), s.org)
        assert client.get("/api/v1/reports/").status_code == 403, role
        assert client.get("/api/v1/reports/contributions/").status_code == 403, role
        assert client.post("/api/v1/reports/people/exports/", {"format": "csv"}, format="json").status_code == 403


def test_reports_are_isolated_per_organization(books, make_org, client_for):
    s = books
    other = make_org("Other Mandal")
    fy = FinancialYear.objects.get(organization=other)
    Event.objects.create(organization=other, financial_year=fy, name="Other Event", code="OE")
    other_admin = client_for(other.admin_user, other)
    # Someone else's event id is simply "not found".
    response = other_admin.get("/api/v1/reports/contributions/", {"event": str(s.event.pk)})
    assert response.status_code == 400 and "event" in response.data["error"]["details"]
    totals = other_admin.get("/api/v1/reports/income-expense/").data["totals"]
    assert totals == {"income": "0.00", "expense": "0.00", "balance": "0.00"}
    # Export jobs belong to the person who asked.
    job = s.admin.post("/api/v1/reports/contributions/exports/", {"format": "csv"}, format="json").data
    assert other_admin.get(f"/api/v1/exports/{job['id']}/").status_code == 404
    assert other_admin.get(f"/api/v1/exports/{job['id']}/download/").status_code == 404


def _download(client, job):
    assert job["status"] == "done", job
    response = client.get(job["download_url"])
    assert response.status_code == 200
    return b"".join(response.streaming_content), response


def test_exports_csv_xlsx_pdf(books):
    s = books
    filters = {"event": str(s.event.pk)}
    job = s.admin.post(
        "/api/v1/reports/income-expense/exports/", {"format": "csv", "filters": filters}, format="json"
    ).data
    content, response = _download(s.admin, job)
    assert content.startswith(b"\xef\xbb\xbf")  # BOM for Excel
    assert "attachment" in response["Content-Disposition"]
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert ["Balance", "636.45"] in rows

    job = s.admin.post(
        "/api/v1/reports/contributions/exports/", {"format": "xlsx", "filters": filters}, format="json"
    ).data
    content, _ = _download(s.admin, job)
    sheet = load_workbook(io.BytesIO(content)).active
    values = [c.value for row in sheet.iter_rows() for c in row]
    assert Decimal("1001.00") in [Decimal(str(v)) for v in values if isinstance(v, int | float | Decimal)]

    job = s.admin.post(
        "/api/v1/reports/expenses/exports/",
        {"format": "pdf", "filters": filters},
        format="json",
        HTTP_ACCEPT_LANGUAGE="mr",
    ).data
    content, _ = _download(s.admin, job)
    assert content.startswith(b"%PDF-")


def test_csv_neutralises_formulas(setup):  # noqa: F811
    s = setup
    Person.objects.create(organization=s.org, full_name='=HYPERLINK("http://evil")')
    job = s.admin.post("/api/v1/reports/people/exports/", {"format": "csv"}, format="json").data
    content, _ = _download(s.admin, job)
    assert "'=HYPERLINK" in content.decode("utf-8-sig")


def test_audit_report_and_export_is_audited(books):
    s = books
    data = s.admin.get("/api/v1/reports/audit/").data
    assert data["row_count"] > 0
    s.admin.post("/api/v1/reports/people/exports/", {"format": "csv"}, format="json")
    actions = {r["action"] for r in s.admin.get("/api/v1/reports/audit/").data["rows"]}
    assert "export" in actions
